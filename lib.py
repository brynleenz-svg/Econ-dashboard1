"""Shared helpers for every country. Needs only Python's standard library + openpyxl.

You normally never edit this file. To add a country, copy scripts/countries/nz.py.
"""
from __future__ import annotations

import calendar
import datetime as dt
import io
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

USER_AGENT = "econ-dashboard/1.0 (open-data reader; +%s)" % os.environ.get(
    "SITE_URL", "https://yoursite.com"
)
# For testing without internet: FIXTURE_DIR=/some/folder python scripts/update.py
FIXTURE_DIR = os.environ.get("FIXTURE_DIR")


def warn(msg: str) -> None:
    print("  warning: " + msg)


# --------------------------------------------------------------------------
# Dates
# --------------------------------------------------------------------------
MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}


def month_end(d: dt.date) -> dt.date:
    return dt.date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def fmt_date(d: dt.date) -> str:
    return f"{d.day} {d.strftime('%b %Y')}"


def parse_date(v) -> dt.date | None:
    """Turn a spreadsheet cell into a date, or None if it isn't one.

    Handles real Excel dates, '12 Oct 2026', 'Sept 2026' (-> end of that month),
    '2026-10-12' and '12/10/2026' (day first, as in NZ).
    """
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    if not isinstance(v, str):
        return None
    s = re.sub(r"\s+", " ", v.strip())
    m = re.fullmatch(r"(\d{1,2}) ([A-Za-z]{3,9})\.? (\d{4})", s)
    if m and m.group(2)[:3].lower() in MONTHS:
        return dt.date(int(m.group(3)), MONTHS[m.group(2)[:3].lower()], int(m.group(1)))
    m = re.fullmatch(r"([A-Za-z]{3,9})\.? (\d{4})", s)
    if m and m.group(1)[:3].lower() in MONTHS:
        return month_end(dt.date(int(m.group(2)), MONTHS[m.group(1)[:3].lower()], 1))
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None


# --------------------------------------------------------------------------
# Downloading
# --------------------------------------------------------------------------
def download(url: str, tries: int = 3) -> bytes:
    last: Exception | None = None
    for attempt in range(tries):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            with urllib.request.urlopen(req, timeout=60) as resp:
                return resp.read()
        except urllib.error.HTTPError as e:
            if e.code < 500:  # 403/404 won't fix themselves, don't hammer the site
                raise RuntimeError(f"HTTP {e.code} for {url}") from e
            last = e
        except Exception as e:  # network blips
            last = e
        time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"Could not download {url}: {last}")


def discover_xlsx(page: str, pick: str, avoid: str | None = None) -> str:
    """Find a spreadsheet link on a web page (used when a file's address isn't fixed)."""
    html = download(page).decode("utf-8", "replace")
    hrefs = re.findall(r'href="([^"]+?\.xlsx[^"]*)"', html, re.I)
    urls = [urllib.parse.urljoin(page, h.replace("&amp;", "&")) for h in hrefs]
    found = [
        u for u in urls
        if re.search(pick, u, re.I) and not (avoid and re.search(avoid, u, re.I))
    ]
    if not found:
        raise RuntimeError(
            f"No spreadsheet link matching /{pick}/ on {page}. Links seen: {urls[:12]}"
        )
    return found[0]


_FILE_CACHE: dict[str, bytes] = {}


def get_file(spec: dict) -> bytes:
    """spec = {'url': ...} and/or {'discover': {page, pick, avoid}}, optional 'fixture' name."""
    key = spec.get("url") or spec["discover"]["page"]
    if key in _FILE_CACHE:
        return _FILE_CACHE[key]
    if FIXTURE_DIR:
        name = spec.get("fixture") or os.path.basename(urllib.parse.urlparse(spec["url"]).path)
        data = (Path(FIXTURE_DIR) / name).read_bytes()
    else:
        data = None
        if spec.get("url"):
            try:
                data = download(spec["url"])
            except RuntimeError as e:
                if "discover" not in spec:
                    raise
                warn(f"{e}; looking for the file on its web page instead")
        if data is None:
            d = spec["discover"]
            data = download(discover_xlsx(d["page"], d["pick"], d.get("avoid")))
    _FILE_CACHE[key] = data
    return data


# --------------------------------------------------------------------------
# Reading spreadsheets without knowing the exact layout
# --------------------------------------------------------------------------
def _clean(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


@dataclass
class Table:
    sheet: str
    labels: list[str]  # one combined header text per column
    rows: list[tuple[dt.date, list]]  # (date in first column, all cells)


def _detect(sheet: str, rows: list[list]) -> Table | None:
    start = -1
    for r in range(min(len(rows) - 1, 80)):
        if parse_date(rows[r][0] if rows[r] else None) and parse_date(
            rows[r + 1][0] if rows[r + 1] else None
        ):
            start = r
            break
    if start < 0:
        return None
    width = max(len(r) for r in rows[: start + 5])
    labels: list[list[str]] = [[] for _ in range(width)]
    for hr in rows[max(0, start - 25): start]:
        filled = sum(1 for c in hr if _clean(c))
        carry = ""
        for c in range(width):
            t = _clean(hr[c]) if c < len(hr) else ""
            if t:
                carry = t
            # Merged group headings ("Cash rate (%pa)") only hold text in their
            # first cell, so spread them rightwards. Single-cell title rows don't spread.
            use = t or (carry if filled >= 2 else "")
            if use and c > 0 and use not in labels[c]:
                labels[c].append(use)
    data = []
    for row in rows[start:]:
        d = parse_date(row[0] if row else None)
        if d:
            data.append((d, row))
    return Table(sheet, [" | ".join(l) for l in labels], data)


def read_tables(data: bytes) -> list[Table]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        try:
            ws.reset_dimensions()
        except Exception:
            pass
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        table = _detect(ws.title, rows)
        if table:
            out.append(table)
    wb.close()
    return out


def _number(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", "").strip())
        except ValueError:
            return None
    return None


def series_from_xlsx(spec: dict, include: list[str], exclude: str | None = None):
    """Return [(date, value), ...] for the first column whose header text matches
    every regex in `include` and not `exclude`."""
    tables = read_tables(get_file(spec))
    seen = []
    for t in tables:
        for c in range(1, len(t.labels)):
            label = t.labels[c]
            seen.append(label)
            if all(re.search(p, label, re.I) for p in include) and not (
                exclude and re.search(exclude, label, re.I)
            ):
                pts = []
                for d, row in t.rows:
                    v = _number(row[c]) if c < len(row) else None
                    if v is not None:
                        pts.append((d, v))
                if len(pts) >= 5:
                    return pts
    raise LookupError(
        f"No column matched {include} (exclude {exclude}). Column headings found: "
        + "; ".join(x for x in seen if x)[:900]
    )


# --------------------------------------------------------------------------
# Turning raw points into what the website needs
# --------------------------------------------------------------------------
def normalise(points, period: str):
    if period == "quarter":
        points = [(month_end(d), v) for d, v in points]
    return sorted({d: v for d, v in points}.items())


def weekly(points):
    last = {}
    for d, v in points:
        last[d.isocalendar()[:2]] = (d, v)
    return sorted(last.values())


def _arrow(x: float) -> str:
    return "▲" if x > 0 else "▼" if x < 0 else "▬"


def describe_change(meta: dict, pts) -> dict:
    spec, dec, last = meta["change"], meta["decimals"], pts[-1][1]
    if spec["type"] == "last_move":
        for i in range(len(pts) - 2, -1, -1):
            if abs(pts[i][1] - last) > 1e-9:
                delta = last - pts[i][1]
                return {
                    "text": f"Last move: {_arrow(delta)} {abs(delta):.{dec}f} pts on {fmt_date(pts[i + 1][0])}",
                    "dir": "up" if delta > 0 else "down",
                }
        return {"text": f"No change since {fmt_date(pts[0][0])}", "dir": "flat"}
    back = spec.get("back", 1)
    if len(pts) <= back:
        return {"text": "", "dir": "flat"}
    prev = pts[-1 - back][1]
    if spec["type"] == "pct":
        delta, unit, places = (last / prev - 1) * 100, "%", 1
    else:
        delta, unit, places = last - prev, " pts", dec
    if round(abs(delta), places) == 0:
        return {"text": f"No change {spec['label']}", "dir": "flat"}
    return {
        "text": f"{_arrow(delta)} {abs(delta):.{places}f}{unit} {spec['label']}",
        "dir": "up" if delta > 0 else "down",
    }


def build_entry(meta: dict, points, sample: bool = False) -> dict:
    pts = normalise(points, meta["period"])
    if len(pts) < 8:
        raise ValueError(f"only {len(pts)} data points found")
    last_d, last_v = pts[-1]
    lo, hi = meta["range"]
    if not lo <= last_v <= hi:
        raise ValueError(
            f"latest value {last_v} is outside the expected range {lo} to {hi}; "
            "the script has probably picked the wrong column"
        )
    hist = pts
    h = meta.get("history", {})
    if h.get("weekly"):
        hist = weekly(hist)
    hist = hist[-h.get("n", 260):]
    places = max(meta["decimals"] + 2, 2)
    age_days = (dt.date.today() - last_d).days
    return {
        "id": meta["id"],
        "label": meta["label"],
        "prefix": meta.get("prefix", ""),
        "suffix": meta.get("suffix", ""),
        "decimals": meta["decimals"],
        "period": meta["period"],
        "shape": meta.get("shape", "line"),
        "blurb": meta.get("blurb", ""),
        "source": {"name": meta["source_name"], "url": meta["source_url"]},
        "value": round(last_v, places),
        "asOf": last_d.isoformat(),
        "change": describe_change(meta, pts),
        "stale": age_days > meta["max_age_days"],
        "sample": sample,
        "series": [[d.isoformat(), round(v, places)] for d, v in hist],
    }
